"""Import USDA FNDDS 2021-2023 and emit reviewed supplementary matches.

The Food and Nutrient Database for Dietary Studies (FNDDS) is part of the
USDA FoodData Central family.  Its nutrient workbook reports values per 100 g
for the named FNDDS food codes.  This importer preserves those codes and the
publisher descriptions, then emits a deliberately small, reviewed mapping
supplement for catalogue records whose state and formulation are represented
by a FNDDS profile.

Usage from the repository root::

    python research/nutrition_coverage/import_fndds.py --download
    python research/nutrition_coverage/import_fndds.py

The raw workbooks are retained under ``research/nutrition_coverage/raw``.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import urllib.request
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook


def _workspace_root() -> Path:
    here = Path(__file__).resolve()
    for candidate in (here.parent, *here.parents):
        if (candidate / ".tmp" / "nutrition_active_catalogue.csv").exists():
            return candidate
    return here.parent


WORKSPACE_ROOT = _workspace_root()
RESEARCH_ROOT = WORKSPACE_ROOT / "research" / "nutrition_coverage"
RAW_DIR = RESEARCH_ROOT / "raw"
OUTPUT_PATH = RESEARCH_ROOT / "fndds_foods.json"
MAPPINGS_PATH = RESEARCH_ROOT / "fndds_mappings.json"
CATALOGUE_PATH = WORKSPACE_ROOT / ".tmp" / "nutrition_active_catalogue.csv"

SOURCE_DOWNLOAD_PAGE = (
    "https://www.ars.usda.gov/northeast-area/beltsville-md-bhnrc/"
    "beltsville-human-nutrition-research-center/food-surveys-research-group/"
    "docs/fndds-download-databases/"
)
NUTRIENT_URL = (
    "https://www.ars.usda.gov/ARSUserFiles/80400530/apps/"
    "2021-2023%20FNDDS%20At%20A%20Glance%20-%20FNDDS%20Nutrient%20Values.xlsx"
)
FOODS_URL = (
    "https://www.ars.usda.gov/ARSUserFiles/80400530/apps/"
    "2021-2023%20FNDDS%20At%20A%20Glance%20-%20Foods%20and%20Beverages.xlsx"
)
NUTRIENT_RAW = RAW_DIR / "fndds_nutrient_values.xlsx"
FOODS_RAW = RAW_DIR / "fndds_foods_beverages.xlsx"

# FNDDS record ids are intentionally explicit here.  The mapping list is kept
# separate from the source import so the parent combiner can accept only the
# entries it wants without altering the existing USDA/AFCD/MyFCD mappings.
INGREDIENT_MATCHES: dict[str, tuple[int, str]] = {
    "1517": (
        63101120,
        "Heinz apple puree is a single-fruit puree; FNDDS unsweetened applesauce matches the fruit and puree state. Age-specific packaging is not represented, and no infant formula profile is used.",
    ),
    "1635": (
        12210400,
        "Nestle Coffee-Mate is a powdered coffee creamer and matches FNDDS generic coffee creamer powder. Brand recipe and fortification can differ.",
    ),
    "1683": (
        71200310,
        "Mister Potato crisps in assorted flavors match FNDDS restructured flavored potato chips; brand flavor blend and seasoning can differ.",
    ),
    "1684": (
        71200310,
        "Pringles assorted flavors match FNDDS restructured flavored potato chips; brand flavor blend and seasoning can differ.",
    ),
    "1685": (
        71200140,
        "Mister Potato chips in assorted flavors match FNDDS regular or kettle-cut sliced potato chips, other flavored; brand flavor blend and seasoning can differ.",
    ),
}

PREPARED_MATCHES: dict[str, tuple[int, str]] = {
    "1203": (
        27115000,
        "Daging masak kicap is beef in soy-based sauce and matches FNDDS beef with soy-based sauce. Cut, oil and seasoning can differ.",
    ),
    "1754": (
        75207021,
        "Sayur taugeh is cooked bean sprouts and matches FNDDS bean sprouts, cooked. Oil and seasoning are not separately specified by the catalogue.",
    ),
    "1756": (
        27120060,
        "Babi masam manis is pork in sweet-and-sour preparation and matches FNDDS sweet-and-sour pork. Cut, sauce and portion can differ.",
    ),
    "1760": (
        26141160,
        "Ikan siakap kukus is steamed sea bass and matches FNDDS bass, steamed. Species naming, fish size, edible portion and seasoning can differ.",
    ),
}


def _download(url: str, destination: Path) -> None:
    destination.parent.mkdir(parents=True, exist_ok=True)
    request = urllib.request.Request(url, headers={"User-Agent": "SmartCart nutrition importer"})
    with urllib.request.urlopen(request, timeout=120) as response:
        destination.write_bytes(response.read())


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _number(value: Any) -> float | None:
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        return None
    result = float(value)
    return result if math.isfinite(result) else None


def _rows(path: Path, sheet_name: str) -> tuple[list[str], Iterable[tuple[Any, ...]]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook[sheet_name]
    iterator = sheet.iter_rows(values_only=True)
    next(iterator, None)  # workbook title row
    headers = [str(value).strip() if value is not None else "" for value in next(iterator)]
    return headers, iterator


def _description(main: Any, additional: Any) -> str:
    base = str(main or "").strip()
    extra = str(additional or "").strip()
    return f"{base}; {extra}" if extra else base


def _load_food_descriptions() -> dict[int, tuple[str, str]]:
    headers, rows = _rows(FOODS_RAW, "Food and Beverages")
    indexes = {name: index for index, name in enumerate(headers)}
    result: dict[int, tuple[str, str]] = {}
    for row in rows:
        try:
            code = int(row[indexes["Food code"]])
        except (KeyError, TypeError, ValueError):
            continue
        result[code] = (
            str(row[indexes["Main food description"]] or "").strip(),
            str(row[indexes["Additional food description"]] or "").strip(),
        )
    return result


def _build_records() -> list[dict[str, Any]]:
    descriptions = _load_food_descriptions()
    headers, rows = _rows(NUTRIENT_RAW, "FNDDS Nutrient Values")
    indexes = {name: index for index, name in enumerate(headers)}
    required = [
        "Food code",
        "Main food description",
        "WWEIA Category number",
        "WWEIA Category description",
        "Energy (kcal)",
        "Protein (g)",
        "Carbohydrate (g)",
        "Sugars, total\n(g)",
        "Fiber, total dietary (g)",
        "Total Fat (g)",
        "Fatty acids, total saturated (g)",
        "Calcium (mg)",
        "Sodium (mg)",
    ]
    missing = [field for field in required if field not in indexes]
    if missing:
        raise ValueError(f"FNDDS nutrient workbook is missing columns: {missing}")

    records: list[dict[str, Any]] = []
    for row in rows:
        try:
            code = int(row[indexes["Food code"]])
        except (TypeError, ValueError):
            continue
        main, additional = descriptions.get(
            code,
            (str(row[indexes["Main food description"]] or "").strip(), ""),
        )
        energy_kcal = _number(row[indexes["Energy (kcal)"]])
        nutrients = {
            "energy_kcal": energy_kcal,
            "energy_kj": round(energy_kcal * 4.184, 4) if energy_kcal is not None else None,
            "protein_g": _number(row[indexes["Protein (g)"]]),
            "fat_g": _number(row[indexes["Total Fat (g)"]]),
            "carbohydrate_g": _number(row[indexes["Carbohydrate (g)"]]),
            "sugars_g": _number(row[indexes["Sugars, total\n(g)"]]),
            "fibre_g": _number(row[indexes["Fiber, total dietary (g)"]]),
            "sodium_mg": _number(row[indexes["Sodium (mg)"]]),
            "saturated_fat_g": _number(row[indexes["Fatty acids, total saturated (g)"]]),
            "calcium_mg": _number(row[indexes["Calcium (mg)"]]),
        }
        records.append(
            {
                "id": f"FNDDS-{code}",
                "source": "USDA-FNDDS-2021-2023",
                "source_code": str(code),
                "source_record_id": str(code),
                "description": _description(main, additional),
                "main_food_description": main,
                "additional_food_description": additional or None,
                "wweia_category_number": int(row[indexes["WWEIA Category number"]]),
                "wweia_category": str(row[indexes["WWEIA Category description"]] or "").strip(),
                "basis": "per_100g",
                "source_url": SOURCE_DOWNLOAD_PAGE,
                "url": SOURCE_DOWNLOAD_PAGE,
                "nutrients": nutrients,
                "complete_minimum": all(
                    nutrients[name] is not None
                    for name in ("energy_kcal", "protein_g", "fat_g", "carbohydrate_g")
                ),
            }
        )
    return records


def build_output(retrieved_at: str | None = None) -> dict[str, Any]:
    records = _build_records()
    timestamp = retrieved_at or datetime.now(timezone.utc).date().isoformat()
    complete_count = sum(1 for record in records if record["complete_minimum"])
    return {
        "schema_version": "1.0",
        "source": "USDA-FNDDS-2021-2023",
        "source_url": SOURCE_DOWNLOAD_PAGE,
        "retrieved_at": timestamp,
        "basis": "per_100g",
        "nutrient_units": {
            "energy_kcal": "kcal",
            "energy_kj": "kJ",
            "protein_g": "g",
            "fat_g": "g",
            "carbohydrate_g": "g",
            "sugars_g": "g",
            "fibre_g": "g",
            "sodium_mg": "mg",
            "saturated_fat_g": "g",
            "calcium_mg": "mg",
        },
        "coverage": {
            "minimum_fields": ["energy_kcal", "protein_g", "fat_g", "carbohydrate_g"],
            "complete_minimum_count": complete_count,
            "nutrient_counts": {
                field: sum(1 for record in records if record["nutrients"][field] is not None)
                for field in (
                    "energy_kcal",
                    "protein_g",
                    "fat_g",
                    "carbohydrate_g",
                    "fibre_g",
                    "sodium_mg",
                    "saturated_fat_g",
                    "calcium_mg",
                )
            },
        },
        "provenance": [
            {
                "publisher": "USDA Agricultural Research Service, Food Surveys Research Group",
                "dataset": "Food and Nutrient Database for Dietary Studies 2021-2023",
                "download_page": SOURCE_DOWNLOAD_PAGE,
                "raw_files": [
                    {
                        "path": f"raw/{NUTRIENT_RAW.name}",
                        "url": NUTRIENT_URL,
                        "sha256": _sha256(NUTRIENT_RAW),
                    },
                    {
                        "path": f"raw/{FOODS_RAW.name}",
                        "url": FOODS_URL,
                        "sha256": _sha256(FOODS_RAW),
                    },
                ],
            }
        ],
        "record_count": len(records),
        "records": records,
        "foods": records,
    }


def _load_catalogue_names() -> dict[str, str]:
    if not CATALOGUE_PATH.exists():
        return {}
    with CATALOGUE_PATH.open(encoding="utf-8-sig", newline="") as stream:
        return {str(row["item_code"]): row["item_name"] for row in csv.DictReader(stream)}


def build_mappings(records: list[dict[str, Any]]) -> list[dict[str, Any]]:
    by_code = {int(record["source_code"]): record for record in records}
    catalogue_names = _load_catalogue_names()
    mappings: list[dict[str, Any]] = []
    for item_code, (food_code, rationale) in (*INGREDIENT_MATCHES.items(), *PREPARED_MATCHES.items()):
        record = by_code[food_code]
        mappings.append(
            {
                "item_code": item_code,
                "item_name": catalogue_names.get(item_code),
                "food_id": record["id"],
                "match_type": "generic",
                "status": "approved",
                "rationale": rationale,
                "source": record["source"],
                "source_record_id": record["source_record_id"],
                "source_description": record["description"],
                "source_url": record["source_url"],
                "basis": record["basis"],
                "nutrients": record["nutrients"],
            }
        )
    return mappings


def write_json(payload: Any, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(path.suffix + ".part")
    temporary.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    temporary.replace(path)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--download", action="store_true", help="download official FNDDS workbooks first")
    parser.add_argument("--output", type=Path, default=OUTPUT_PATH)
    parser.add_argument("--mappings-output", type=Path, default=MAPPINGS_PATH)
    args = parser.parse_args()
    if args.download:
        _download(NUTRIENT_URL, NUTRIENT_RAW)
        _download(FOODS_URL, FOODS_RAW)
    for path in (NUTRIENT_RAW, FOODS_RAW):
        if not path.exists():
            raise FileNotFoundError(f"Missing {path}; run with --download")
    payload = build_output()
    mappings = build_mappings(payload["records"])
    write_json(payload, args.output)
    write_json(mappings, args.mappings_output)
    print(
        f"Wrote {payload['record_count']} FNDDS records to {args.output} "
        f"({payload['coverage']['complete_minimum_count']} complete minimum records)"
    )
    print(f"Wrote {len(mappings)} reviewed supplementary mappings to {args.mappings_output}")


if __name__ == "__main__":
    main()
